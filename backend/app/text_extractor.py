import os
import logging
from pathlib import Path
from typing import Optional
from app.config import settings

logger = logging.getLogger("folderpilot.extractor")

# Conditional imports with graceful degradation
try:
    import fitz  # PyMuPDF
    HAVE_PYMUPDF = True
except ImportError:
    HAVE_PYMUPDF = False

try:
    import docx
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False

try:
    import pptx
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False

try:
    from rapidocr_onnxruntime import RapidOCR
    rapid_ocr = RapidOCR()
    HAVE_RAPIDOCR = True
except Exception:
    rapid_ocr = None
    HAVE_RAPIDOCR = False

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".yaml", ".yml",
    ".xml", ".html", ".htm", ".log", ".ini", ".cfg", ".conf", ".py",
    ".js", ".jsx", ".ts", ".tsx", ".css", ".scss", ".java", ".c", ".cpp",
    ".h", ".sql", ".sh", ".bat", ".ps1"
}

def extract_text_from_file(filepath: Path) -> Optional[str]:
    """
    Safely extracts the first ~500 tokens of text from supported file types.
    Never raises an uncaught exception (FR-14).
    """
    if not filepath.exists() or filepath.is_dir():
        return None
        
    try:
        size = filepath.stat().st_size
        if size > settings.MAX_FILE_SIZE_EXTRACT or size == 0:
            return None
    except OSError:
        return None

    ext = filepath.suffix.lower()
    extracted = ""

    try:
        # Plain text and code files
        if ext in TEXT_EXTENSIONS:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                extracted = f.read(settings.MAX_TEXT_SNIPPET_CHARS)
                
        # PDF documents
        elif ext == ".pdf" and HAVE_PYMUPDF:
            with fitz.open(filepath) as doc:
                text_parts = []
                for page_num in range(min(3, len(doc))): # First 3 pages max
                    page = doc[page_num]
                    text_parts.append(page.get_text())
                    if sum(len(p) for p in text_parts) >= settings.MAX_TEXT_SNIPPET_CHARS:
                        break
                extracted = " ".join(text_parts)[:settings.MAX_TEXT_SNIPPET_CHARS]
                
        # DOCX documents
        elif ext == ".docx" and HAVE_DOCX:
            doc = docx.Document(filepath)
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            extracted = " ".join(paragraphs)[:settings.MAX_TEXT_SNIPPET_CHARS]
            
        # PPTX presentations
        elif ext == ".pptx" and HAVE_PPTX:
            prs = pptx.Presentation(filepath)
            slides_text = []
            for slide in prs.slides[:5]: # First 5 slides max
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            slides_text.append(paragraph.text)
            extracted = " ".join(slides_text)[:settings.MAX_TEXT_SNIPPET_CHARS]
            
        # Images (RapidOCR if available and under 10MB)
        elif ext in {".png", ".jpg", ".jpeg", ".bmp", ".webp"} and HAVE_RAPIDOCR and size < 10 * 1024 * 1024:
            result, _ = rapid_ocr(str(filepath))
            if result:
                text_lines = [line[1] for line in result if line and len(line) > 1]
                extracted = " ".join(text_lines)[:settings.MAX_TEXT_SNIPPET_CHARS]

    except Exception as e:
        logger.debug(f"Could not extract text from {filepath.name}: {e}")
        return None

    cleaned = " ".join(extracted.split())
    return cleaned if cleaned else None
