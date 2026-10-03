import os
import uuid
import asyncio
import logging
from pathlib import Path
from typing import Optional, List, Dict, Any
from contextlib import asynccontextmanager
import mimetypes
import subprocess
from fastapi import FastAPI, HTTPException, BackgroundTasks, Query, Depends, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from pydantic import BaseModel

from app.config import settings
from app.database import init_db, get_db, query_all, query_one
from app.folder_browser import browse_directory, is_path_protected, open_native_folder_dialog
from app.scanner import scan_workspace_worker, scan_jobs
from app.chaos_score import compute_chaos_score
from app.planner import generate_plan, calculate_dry_run, apply_plan
from app.journal import JournalEngine
from app.chat_engine import handle_chat_message

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
logger = logging.getLogger("folderpilot.api")

@asynccontextmanager
async def lifespan(app: FastAPI):
    # Startup: ensure SQLite database and tables exist
    init_db()
    logger.info("FolderPilot backend ready on 127.0.0.1:8000")
    yield

app = FastAPI(
    title=settings.APP_NAME,
    version=settings.VERSION,
    lifespan=lifespan
)

# CORS setup for localhost development and preview
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- Request / Response Models ---

class CreateWorkspaceRequest(BaseModel):
    path: str
    read_only: bool = False

class ScanRequest(BaseModel):
    recursive: bool = True
    max_depth: int = 10

class ChatRequest(BaseModel):
    message: str

class UpdateOpsRequest(BaseModel):
    op_ids: List[str]
    approved: bool

class OverrideClassificationRequest(BaseModel):
    category: str
    subfolder: Optional[str] = None

class UndoRequest(BaseModel):
    scope: str = "batch" # 'batch', 'all', or 'op'
    target_id: Optional[str] = None # plan_id, workspace_id, or journal_id

# --- API Endpoints ---

@app.get("/api/health")
def health_check():
    return {"status": "ok", "app": settings.APP_NAME, "version": settings.VERSION}

@app.get("/fs/browse")
def browse_fs(path: Optional[str] = Query(None)):
    """Server-side folder browser (FR-01, FR-02)."""
    return browse_directory(path)

@app.api_route("/fs/browse-dialog", methods=["GET", "POST"])
def browse_fs_dialog(initial_path: Optional[str] = Query(None)):
    """Opens native Windows folder picker dialog (like in VS Code) and returns selected path."""
    selected = open_native_folder_dialog(initial_path)
    if selected:
        return {"status": "selected", "path": selected}
    return {"status": "cancelled", "path": None}


@app.api_route("/fs/file", methods=["GET", "HEAD"])
def get_fs_file(path: str = Query(...)):
    """Streams file for in-app preview (images, audio, video, pdf). Supports HEAD probing and byte ranges."""
    target = Path(path).resolve()
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    if is_path_protected(target):
        raise HTTPException(status_code=403, detail="Access to protected system file denied")
    
    media_type, _ = mimetypes.guess_type(str(target))
    headers = {
        "Accept-Ranges": "bytes",
        "Cache-Control": "public, max-age=3600"
    }
    return FileResponse(
        path=str(target), 
        media_type=media_type or "application/octet-stream", 
        filename=target.name,
        headers=headers,
        content_disposition_type="inline"
    )

@app.get("/fs/file/text")
def get_fs_file_text(path: str = Query(...), max_bytes: int = Query(512000)):
    """Reads text/code/document file snippet for syntax-highlighted and formatted in-app preview."""
    target = Path(path).resolve()
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    if is_path_protected(target):
        raise HTTPException(status_code=403, detail="Access to protected system file denied")

    file_size = target.stat().st_size
    ext = target.suffix.lower()
    truncated = False
    content = ""
    is_binary = False
    doc_type = "text"

    try:
        # 1. DOCX Documents
        if ext == ".docx":
            doc_type = "docx"
            try:
                import docx
                doc = docx.Document(str(target))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                content = "\n\n".join(paragraphs)
                if not content:
                    content = "[Document contains no text paragraphs or tables]"
            except Exception as e:
                content = f"// Error reading DOCX: {e}"

        # 2. PPTX Presentations
        elif ext == ".pptx":
            doc_type = "pptx"
            try:
                import pptx
                prs = pptx.Presentation(str(target))
                slide_texts = []
                for idx, slide in enumerate(prs.slides[:20], 1):
                    lines = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for p in shape.text_frame.paragraphs:
                                if p.text.strip():
                                    lines.append(p.text.strip())
                    if lines:
                        slide_texts.append(f"--- Slide {idx} ---\n" + "\n".join(lines))
                content = "\n\n".join(slide_texts) or "[Presentation contains no text on first 20 slides]"
            except Exception as e:
                content = f"// Error reading PPTX: {e}"

        # 3. PDF Documents
        elif ext == ".pdf":
            doc_type = "pdf"
            try:
                import fitz
                with fitz.open(str(target)) as doc:
                    page_texts = []
                    for page_num in range(min(10, len(doc))):
                        t = doc[page_num].get_text()
                        if t.strip():
                            page_texts.append(f"--- Page {page_num + 1} of {len(doc)} ---\n" + t.strip())
                    content = "\n\n".join(page_texts) or "[PDF contains no selectable text (scanned image or empty)]"
            except Exception as e:
                content = f"// Error reading PDF text: {e}"

        # 4. Plain Text and Code
        else:
            truncated = file_size > max_bytes
            try:
                with open(target, "r", encoding="utf-8", errors="strict") as f:
                    content = f.read(max_bytes)
            except (UnicodeDecodeError, ValueError):
                # Try latin-1 or detect if binary
                try:
                    with open(target, "rb") as f:
                        raw = f.read(min(file_size, 1024))
                    if b"\x00" in raw:
                        # Definitely binary
                        is_binary = True
                        doc_type = "binary"
                        # Generate clean hex dump preview
                        hex_lines = []
                        chunk = raw[:512]
                        for i in range(0, len(chunk), 16):
                            b16 = chunk[i:i+16]
                            hex_str = " ".join(f"{b:02x}" for b in b16)
                            ascii_str = "".join(chr(b) if 32 <= b < 127 else "." for b in b16)
                            hex_lines.append(f"{i:08x}  {hex_str:<48}  |{ascii_str}|")
                        content = "\n".join(hex_lines)
                    else:
                        with open(target, "r", encoding="latin-1", errors="replace") as f:
                            content = f.read(max_bytes)
                except Exception as e:
                    content = f"// Binary file ({file_size} bytes). Cannot preview text: {e}"
                    is_binary = True
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Could not read file text: {str(e)}")

    word_count = len(content.split()) if not is_binary else 0

    return {
        "name": target.name,
        "extension": ext,
        "size": file_size,
        "truncated": truncated,
        "content": content,
        "is_binary": is_binary,
        "doc_type": doc_type,
        "word_count": word_count
    }

class RevealRequest(BaseModel):
    path: str

@app.post("/fs/reveal")
def reveal_in_explorer(req: RevealRequest):
    """Reveals the file or folder in Windows File Explorer."""
    target = Path(req.path).resolve()
    if not target.exists():
        raise HTTPException(status_code=404, detail="File or folder not found")
    try:
        if os.name == 'nt':
            if target.is_file():
                subprocess.Popen(f'explorer /select,"{target}"')
            else:
                subprocess.Popen(f'explorer "{target}"')
            return {"status": "revealed", "path": str(target)}
        else:
            return {"status": "unsupported_os", "path": str(target)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Could not open explorer: {str(e)}")

@app.post("/workspaces")
def create_workspace(req: CreateWorkspaceRequest):
    """Creates or connects to an existing workspace (FR-01, FR-02, FR-06)."""
    target = Path(req.path).resolve()
    if not target.exists() or not target.is_dir():
        raise HTTPException(status_code=400, detail=f"Directory '{target}' does not exist.")
    if is_path_protected(target):
        raise HTTPException(
            status_code=403, 
            detail="Protected system folder cannot be selected as a workspace."
        )

    existing = query_one("SELECT * FROM workspaces WHERE root_path = ?", (str(target),))
    if existing:
        return existing

    ws_id = str(uuid.uuid4())
    with get_db() as conn:
        conn.execute(
            "INSERT INTO workspaces (id, root_path, read_only) VALUES (?, ?, ?)",
            (ws_id, str(target), int(req.read_only))
        )
    return query_one("SELECT * FROM workspaces WHERE id = ?", (ws_id,))

@app.get("/workspaces")
def list_workspaces():
    return query_all("SELECT * FROM workspaces ORDER BY created_at DESC")

@app.get("/workspaces/{ws_id}")
def get_workspace(ws_id: str):
    ws = query_one("SELECT * FROM workspaces WHERE id = ?", (ws_id,))
    if not ws:
        raise HTTPException(status_code=404, detail="Workspace not found")
    return ws

@app.post("/workspaces/{ws_id}/scan")
def start_scan(ws_id: str, req: ScanRequest, bg_tasks: BackgroundTasks):
    """Triggers background scanning job (FR-10)."""
    ws = query_one("SELECT * FROM workspaces WHERE id = ?", (ws_id,))
    if not ws:
        raise HTTPException(status_code=404, detail="Workspace not found")

    job_id = str(uuid.uuid4())
    scan_jobs[job_id] = {
        "job_id": job_id,
        "workspace_id": ws_id,
        "status": "pending",
        "files_seen": 0,
        "total_files": 0,
        "bytes_seen": 0,
        "current_file": "Initializing scan..."
    }

    bg_tasks.add_task(scan_workspace_worker, ws_id, job_id, req.recursive, req.max_depth)
    return {"job_id": job_id, "status": "started"}

@app.get("/jobs/{job_id}")
def get_job_status(job_id: str):
    """Poll or check status of a background job."""
    if job_id not in scan_jobs:
        raise HTTPException(status_code=404, detail="Job not found")
    return scan_jobs[job_id]

@app.get("/jobs/{job_id}/stream")
async def stream_job_progress(job_id: str):
    """Server-Sent Events (SSE) live progress stream (FR-10)."""
    if job_id not in scan_jobs:
        raise HTTPException(status_code=404, detail="Job not found")

    async def event_generator():
        while True:
            job = scan_jobs.get(job_id)
            if not job:
                break
            data = (
                f"data: {{\"status\": \"{job['status']}\", \"files_seen\": {job['files_seen']}, "
                f"\"total_files\": {job.get('total_files', 0)}, \"bytes_seen\": {job['bytes_seen']}, "
                f"\"current_file\": \"{job.get('current_file', '')}\"}}\n\n"
            )
            yield data
            if job["status"] in ("completed", "failed", "cancelled"):
                break
            await asyncio.sleep(0.4)

    return StreamingResponse(event_generator(), media_type="text/event-stream")

@app.get("/workspaces/{ws_id}/stats")
def get_workspace_stats(ws_id: str):
    """Overview dashboard stats and Chaos Score (FR-40, FR-48)."""
    ws = query_one("SELECT * FROM workspaces WHERE id = ?", (ws_id,))
    if not ws:
        raise HTTPException(status_code=404, detail="Workspace not found")

    files = query_all("SELECT * FROM files WHERE workspace_id = ?", (ws_id,))
    total_files = len(files)
    total_bytes = sum(f["size"] for f in files)

    # Category counts
    cat_counts = query_all(
        """
        SELECT c.category, COUNT(*) as count, SUM(f.size) as total_size
        FROM classifications c
        JOIN files f ON c.file_id = f.id
        WHERE f.workspace_id = ?
        GROUP BY c.category
        """,
        (ws_id,)
    )

    # Duplicates count
    dupes_row = query_one(
        """
        SELECT COUNT(dm.file_id) as cnt
        FROM dup_members dm
        JOIN dup_groups dg ON dm.group_id = dg.id
        WHERE dg.workspace_id = ?
        """,
        (ws_id,)
    )
    dup_count = dupes_row["cnt"] if dupes_row else 0

    # Sensitive files count
    sensitive_row = query_one(
        """
        SELECT COUNT(*) as cnt
        FROM classifications c
        JOIN files f ON c.file_id = f.id
        WHERE f.workspace_id = ? AND c.is_sensitive = 1
        """,
        (ws_id,)
    )
    sensitive_count = sensitive_row["cnt"] if sensitive_row else 0

    chaos = compute_chaos_score(files, ws["root_path"], duplicates_count=dup_count)

    return {
        "workspace_id": ws_id,
        "root_path": ws["root_path"],
        "total_files": total_files,
        "total_bytes": total_bytes,
        "category_distribution": cat_counts,
        "duplicate_files_count": dup_count,
        "sensitive_files_count": sensitive_count,
        "chaos_score": chaos
    }

@app.get("/workspaces/{ws_id}/tree")
def get_workspace_tree(ws_id: str, view: str = Query("current")):
    """
    Returns hierarchical JSON tree for D3 Treemap (FR-41) and Before/After Tree (FR-42).
    """
    ws = query_one("SELECT * FROM workspaces WHERE id = ?", (ws_id,))
    if not ws:
        raise HTTPException(status_code=404, detail="Workspace not found")

    files = query_all(
        """
        SELECT f.id, f.path, f.name, f.ext, f.size, f.mtime,
               c.category, c.subfolder, c.confidence, c.reason, c.is_sensitive
        FROM files f
        LEFT JOIN classifications c ON f.id = c.file_id
        WHERE f.workspace_id = ?
        """,
        (ws_id,)
    )

    root_path = Path(ws["root_path"])

    if view == "proposed":
        # Group files under their suggested category folder
        tree: Dict[str, Any] = {"name": root_path.name, "children": {}, "is_dir": True}
        for f in files:
            cat = f.get("category") or "Others"
            sub = f.get("subfolder") or cat
            if cat not in tree["children"]:
                tree["children"][cat] = {"name": cat, "children": {}, "is_dir": True, "category": cat}
            
            tree["children"][cat]["children"][f["name"]] = {
                "id": f["id"],
                "path": f["path"],
                "name": f["name"],
                "size": f["size"],
                "category": cat,
                "confidence": f.get("confidence", 0.8),
                "reason": f.get("reason", ""),
                "is_sensitive": bool(f.get("is_sensitive")),
                "is_dir": False
            }
        
        # Convert nested dicts to array of children for D3 hierarchy
        def format_node(d):
            children = []
            for k, v in d.get("children", {}).items():
                if v.get("is_dir"):
                    children.append(format_node(v))
                else:
                    children.append(v)
            return {"name": d["name"], "category": d.get("category"), "children": children}

        return format_node(tree)

    else:
        # Current physical tree
        tree_root: Dict[str, Any] = {"name": root_path.name, "children": {}, "is_dir": True}
        for f in files:
            p = Path(f["path"])
            try:
                rel_parts = p.relative_to(root_path).parts
            except ValueError:
                rel_parts = [p.name]

            curr = tree_root
            for part in rel_parts[:-1]:
                curr = curr["children"].setdefault(part, {"name": part, "children": {}, "is_dir": True})
            
            curr["children"][p.name] = {
                "id": f["id"],
                "path": f["path"],
                "name": f["name"],
                "size": f["size"],
                "category": f.get("category", "Others"),
                "confidence": f.get("confidence", 0.8),
                "reason": f.get("reason", ""),
                "is_sensitive": bool(f.get("is_sensitive")),
                "is_dir": False
            }

        def format_current(d):
            children = []
            for k, v in d.get("children", {}).items():
                if v.get("is_dir"):
                    children.append(format_current(v))
                else:
                    children.append(v)
            return {"name": d["name"], "children": children}

        return format_current(tree_root)

@app.get("/workspaces/{ws_id}/files")
def get_files(
    ws_id: str, 
    category: Optional[str] = None, 
    search: Optional[str] = None,
    sensitive_only: bool = False
):
    query = """
        SELECT f.*, c.category, c.subfolder, c.confidence, c.tier, c.reason, c.is_sensitive
        FROM files f
        LEFT JOIN classifications c ON f.id = c.file_id
        WHERE f.workspace_id = ?
    """
    params = [ws_id]
    if category:
        query += " AND c.category = ?"
        params.append(category)
    if search:
        query += " AND (f.name LIKE ? OR f.text_snippet LIKE ?)"
        params.extend([f"%{search}%", f"%{search}%"])
    if sensitive_only:
        query += " AND c.is_sensitive = 1"
        
    query += " ORDER BY f.name ASC LIMIT 500"
    return query_all(query, tuple(params))

@app.get("/workspaces/{ws_id}/duplicates")
def get_duplicates(ws_id: str):
    """Returns clustered duplicate groups with recommended keep file (FR-46)."""
    groups = query_all("SELECT * FROM dup_groups WHERE workspace_id = ?", (ws_id,))
    res = []
    for g in groups:
        members = query_all(
            """
            SELECT f.*, c.category
            FROM dup_members dm
            JOIN files f ON dm.file_id = f.id
            LEFT JOIN classifications c ON f.id = c.file_id
            WHERE dm.group_id = ?
            ORDER BY f.mtime DESC
            """,
            (g["id"],)
        )
        res.append({
            "group_id": g["id"],
            "kind": g["kind"],
            "keep_file_id": g["keep_file_id"],
            "members": members
        })
    return res

@app.patch("/classification/{file_id}")
def override_classification(file_id: str, req: OverrideClassificationRequest):
    """Interactive customization: user overrides category (FR-43, FR-27)."""
    file_record = query_one("SELECT * FROM files WHERE id = ?", (file_id,))
    if not file_record:
        raise HTTPException(status_code=404, detail="File not found")

    old_cls = query_one("SELECT category FROM classifications WHERE file_id = ?", (file_id,))
    old_cat = old_cls["category"] if old_cls else "Others"

    with get_db() as conn:
        conn.execute(
            """
            UPDATE classifications
            SET category = ?, subfolder = ?, user_override = 1, confidence = 1.0, reason = 'User manual override'
            WHERE file_id = ?
            """,
            (req.category, req.subfolder, file_id)
        )
        # Store for active learning (FR-27)
        conn.execute(
            "INSERT INTO corrections (id, workspace_id, file_id, from_category, to_category) VALUES (?, ?, ?, ?, ?)",
            (str(uuid.uuid4()), file_record["workspace_id"], file_id, file_record["name"], req.category)
        )
    return {"status": "updated", "file_id": file_id, "new_category": req.category}

@app.post("/workspaces/{ws_id}/plan")
def create_plan(ws_id: str):
    """Generates draft organization plan (FR-50)."""
    return generate_plan(ws_id)

@app.get("/workspaces/{ws_id}/plans")
def list_plans(ws_id: str):
    return query_all("SELECT * FROM plans WHERE workspace_id = ? ORDER BY created_at DESC", (ws_id,))

@app.get("/plan/{plan_id}/ops")
def get_plan_ops(plan_id: str):
    return query_all("SELECT * FROM ops WHERE plan_id = ?", (plan_id,))

@app.patch("/plan/{plan_id}/ops")
def update_plan_ops(plan_id: str, req: UpdateOpsRequest):
    """Approves or unapproves operations (FR-51)."""
    with get_db() as conn:
        for op_id in req.op_ids:
            conn.execute("UPDATE ops SET approved = ? WHERE id = ? AND plan_id = ?", (int(req.approved), op_id, plan_id))
    return {"status": "updated", "affected": len(req.op_ids), "approved": req.approved}

@app.get("/plan/{plan_id}/dryrun")
def get_dry_run(plan_id: str):
    """Dry run summary and conflict checks (FR-52)."""
    return calculate_dry_run(plan_id)

@app.post("/plan/{plan_id}/apply")
def execute_plan(plan_id: str):
    """Applies approved operations atomically with journal (FR-53, FR-56)."""
    return apply_plan(plan_id)

@app.get("/workspaces/{ws_id}/journal")
def get_journal(ws_id: str):
    """Returns append-only journal entries for undo history (FR-56, FR-57)."""
    return query_all("SELECT * FROM journal WHERE workspace_id = ? ORDER BY ts DESC", (ws_id,))

@app.post("/journal/undo")
def trigger_undo(req: UndoRequest):
    """Undoes operations: batch, single op, or all workspace ops (FR-57)."""
    if req.scope == "batch" and req.target_id:
        plan = query_one("SELECT * FROM plans WHERE id = ?", (req.target_id,))
        if not plan:
            raise HTTPException(status_code=404, detail="Plan not found")
        ws = query_one("SELECT * FROM workspaces WHERE id = ?", (plan["workspace_id"],))
        return JournalEngine.undo_batch(req.target_id, Path(ws["root_path"]))

    elif req.scope == "all" and req.target_id:
        ws = query_one("SELECT * FROM workspaces WHERE id = ?", (req.target_id,))
        if not ws:
            raise HTTPException(status_code=404, detail="Workspace not found")
        return JournalEngine.undo_all_for_workspace(req.target_id, Path(ws["root_path"]))

    elif req.scope == "op" and req.target_id:
        entry = query_one("SELECT * FROM journal WHERE id = ?", (req.target_id,))
        if not entry:
            raise HTTPException(status_code=404, detail="Journal entry not found")
        ws = query_one("SELECT * FROM workspaces WHERE id = ?", (entry["workspace_id"],))
        return JournalEngine.undo_operation(entry, Path(ws["root_path"]))

    raise HTTPException(status_code=400, detail="Invalid undo parameters.")

@app.post("/workspaces/{ws_id}/chat")
def chat_with_folder(ws_id: str, req: ChatRequest):
    """Deterministic chat engine with database queries and plan generation (FR-70 - FR-80)."""
    ws = query_one("SELECT * FROM workspaces WHERE id = ?", (ws_id,))
    if not ws:
        raise HTTPException(status_code=404, detail="Workspace not found")

    response = handle_chat_message(ws_id, req.message)
    # Save chat history
    with get_db() as conn:
        conn.execute(
            "INSERT INTO chat_messages (id, workspace_id, role, content) VALUES (?, ?, 'user', ?)",
            (str(uuid.uuid4()), ws_id, req.message)
        )
        conn.execute(
            "INSERT INTO chat_messages (id, workspace_id, role, content) VALUES (?, ?, 'assistant', ?)",
            (str(uuid.uuid4()), ws_id, response["content"])
        )
    return response

@app.get("/workspaces/{ws_id}/chat/history")
def get_chat_history(ws_id: str):
    return query_all("SELECT * FROM chat_messages WHERE workspace_id = ? ORDER BY ts ASC", (ws_id,))

# --- Local Ollama AI Endpoints ---

class SetModelRequest(BaseModel):
    model: str

class SummarizeRequest(BaseModel):
    path: Optional[str] = None
    filename: Optional[str] = None
    content: Optional[str] = None

@app.get("/api/ai/status")
def get_ai_status():
    """Returns local Ollama service status, active model, and installed models."""
    from app.ollama_client import OllamaClient
    return OllamaClient.get_status()

@app.post("/api/ai/model")
def set_ai_model(req: SetModelRequest):
    """Sets active model for OllamaClient."""
    from app.ollama_client import OllamaClient
    OllamaClient.set_model(req.model)
    return OllamaClient.get_status()

@app.post("/api/ai/summarize")
def summarize_file_with_ai(req: SummarizeRequest):
    """Summarizes a file or snippet using Qwen 2.5 local model."""
    from app.ollama_client import OllamaClient
    from app.scanner import extract_text_from_file

    if not OllamaClient.is_available():
        raise HTTPException(status_code=503, detail="Local Ollama AI service is not running or unreachable at 127.0.0.1:11434.")

    filename = req.filename or "Document"
    content = req.content or ""

    if req.path:
        p = Path(req.path).resolve()
        if not p.exists() or not p.is_file():
            raise HTTPException(status_code=404, detail="File not found")
        filename = p.name
        if not content:
            content = extract_text_from_file(p) or ""

    if not content.strip():
        raise HTTPException(status_code=400, detail="No readable text available to summarize.")

    res = OllamaClient.summarize_content(filename, content)
    if not res:
        raise HTTPException(status_code=500, detail="AI summarization failed or timed out.")
    return res

